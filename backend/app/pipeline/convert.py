"""Deterministic conversion of immutable attachments into canonical shipping text."""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import logging
import os
import os
import re
import shutil
import subprocess
import time as clock
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor
from typing import Any, Iterable

import pymupdf
from pypdf import PdfReader
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from ..config import Settings, settings_from_env
from ..db import database, initialize
from .ingest import attachment_path, ingest


CANONICAL_HEADINGS = {"SI": "SHIPPING INSTRUCTION", "BL": "BILL OF LADING (DRAFT)"}
SEPARATOR = "=" * 40
MIN_LABEL_ROWS = 3
SEVEN_FIELDS = {
    "shipper": ("shipper", "exporter", "consignor"),
    "consignee": ("consignee", "to the order of", "consigned to"),
    "notify_party": ("notify",),
    "port_of_loading": ("port of loading", "loading port", "load port", "pol"),
    "port_of_discharge": ("port of discharge", "discharge port", "destination port", "pod"),
    "container_count": ("container", "containers"),
    "gross_weight_kg": ("gross weight", "gross wt", "total gross"),
}
LAYOUT_FIELD = re.compile(r"^(?P<label>\S.*?)\s{2,}(?P<value>\S.*)$")
GEMINI_PROMPT_VERSION = "scan-extraction-v1"
GEMINI_SCAN_PROMPT = """Copy the shipping details exactly as printed. Return JSON only with raw_lines, fields (shipper, consignee, notify_party, port_of_loading, port_of_discharge, container_count, gross_weight_kg), and uncertain. Do not correct or guess."""
_gemini_last_request = 0.0


class ConversionError(Exception):
    """Expected failure caused by an untrusted source attachment."""


@dataclass(frozen=True)
class FieldLine:
    label: str
    value: str


@dataclass(frozen=True)
class Conversion:
    doc_id: str
    text: str
    metadata: dict[str, Any]


def _plain(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")


def _role_from_name(document: dict[str, Any]) -> tuple[str | None, list[str]]:
    stem = Path(document["path"]).stem.upper()
    if stem.endswith("_SI"):
        return "SI", ["filename_suffix=_SI"]
    if stem.endswith("_BL"):
        return "BL", ["filename_suffix=_BL"]
    return None, []


def _role_from_text(text: str) -> tuple[str | None, str | None]:
    normal = text.upper()
    if "SHIPPING INSTRUCTION" in normal or "BILL OF LADING INSTRUCTION" in normal:
        return "SI", "content=shipping_instruction"
    if "BILL OF LADING" in normal or "B/L NO." in normal:
        return "BL", "content=bill_of_lading"
    return None, None


def _resolve_role(document: dict[str, Any], evidence: list[tuple[str | None, str | None]]) -> tuple[str, list[str], list[str]]:
    filename_role, filename_evidence = _role_from_name(document)
    candidates: list[tuple[str, str]] = []
    if filename_role:
        candidates.extend((filename_role, item) for item in filename_evidence)
    candidates.extend((role, detail) for role, detail in evidence if role and detail)
    if not candidates:
        raise ConversionError("missing_role_hint")
    role = filename_role or candidates[0][0]
    role_evidence = [detail for _, detail in candidates]
    conflicts = sorted({candidate for candidate, _ in candidates if candidate != role})
    warnings = [f"wrong_doc_type_candidate: expected={role}, observed={other}" for other in conflicts]
    return role, role_evidence, warnings


def _field_line(label: str, value: str) -> list[str]:
    label = _plain(label).strip()
    parts = [part.strip() for part in _plain(value).split("\n") if part.strip()]
    first = parts[0] if parts else ""
    separator = "" if label.endswith(":") else ":"
    line = f"{label}{separator}" if not first else f"{label}{separator} {first}"
    if len(parts) > 1:
        line += "\n  " + "; ".join(parts[1:])
    return [line]


def canonical_text(role: str, unlabelled: Iterable[str], fields: Iterable[FieldLine]) -> str:
    lines = [CANONICAL_HEADINGS[role], SEPARATOR, ""]
    for value in unlabelled:
        value = _plain(value).strip()
        if value:
            lines.append(value)
    if any(_plain(value).strip() for value in unlabelled):
        lines.append("")
    for field in fields:
        lines.extend(_field_line(field.label, field.value))
    return "\n".join(lines).rstrip() + "\n"


def canonical_text_from_items(role: str, items: Iterable[str | FieldLine]) -> str:
    """Write source blocks in their original order, including tables between paragraphs."""
    lines = [CANONICAL_HEADINGS[role], SEPARATOR, ""]
    for item in items:
        if isinstance(item, FieldLine):
            lines.extend(_field_line(item.label, item.value))
        else:
            value = _plain(item).strip()
            if value:
                lines.append(value)
    return "\n".join(lines).rstrip() + "\n"


def _without_redundant_role_title(items: list[str | FieldLine], role: str) -> list[str | FieldLine]:
    """The role-derived canonical heading already retains an identical Word title."""
    for index, item in enumerate(items):
        if isinstance(item, FieldLine):
            break
        if _plain(item).strip() == CANONICAL_HEADINGS[role]:
            return [*items[:index], *items[index + 1 :]]
    return items


def canonical_fields(text: str) -> dict[str, str]:
    """Small, non-normalizing canonical parser reused by the Phase 4 quality report."""
    found: dict[str, str] = {}
    for line in text.splitlines():
        label, separator, value = line.partition(":")
        if not separator:
            continue
        normalized = label.lower()
        for field, aliases in SEVEN_FIELDS.items():
            if field not in found and any(alias in normalized for alias in aliases):
                found[field] = value.strip()
    return found


def validate_canonical(text: str, *, require_shipping_fields: bool = True) -> list[str]:
    lines = text.splitlines()
    errors: list[str] = []
    if not lines or lines[0] not in CANONICAL_HEADINGS.values():
        errors.append("missing_or_unknown_heading")
    if len(lines) < 2 or lines[1] != SEPARATOR:
        errors.append("missing_separator")
    label_rows = [line for line in lines[3:] if ":" in line and not line.startswith("  ")]
    if len(label_rows) < MIN_LABEL_ROWS:
        errors.append(f"too_few_label_value_lines:{len(label_rows)}")
    if require_shipping_fields and len(canonical_fields(text)) < 2:
        errors.append("too_few_recognized_shipping_fields")
    return errors


def _xlsx_records(source: Path, document: dict[str, Any]) -> tuple[list[str], list[FieldLine], list[str], list[tuple[str | None, str | None]]]:
    warnings: list[str] = []
    workbook = load_workbook(source, read_only=False, data_only=True)
    try:
        active = workbook.active
        matching = [sheet for sheet in workbook.worksheets if _sheet_role(sheet.title)]
        filename_role, _ = _role_from_name(document)
        selected = next((sheet for sheet in matching if _sheet_role(sheet.title) == filename_role), active)
        if len(workbook.sheetnames) > 1:
            warnings.append(f"extra_sheets:{','.join(name for name in workbook.sheetnames if name != selected.title)}")
        if selected.title != active.title:
            warnings.append(f"selected_role_sheet:{selected.title}")
        if selected.merged_cells.ranges:
            warnings.append(f"merged_cells:{len(selected.merged_cells.ranges)}")
        unlabelled: list[str] = []
        fields: list[FieldLine] = []
        for row_index, cells in enumerate(selected.iter_rows(min_col=1, max_col=2), start=1):
            label = _plain(cells[0].value)
            value = _plain(cells[1].value)
            if not label.strip() and not value.strip():
                continue
            if row_index == 1 and label.strip() and not value.strip():
                unlabelled.append(label)
                continue
            if not label.strip():
                warnings.append(f"blank_label_row:{row_index}")
                unlabelled.append(value)
                continue
            if " | " in value:
                name, address = value.split(" | ", 1)
                value = name + "\n" + address
            fields.append(FieldLine(label, value))
        evidence = [(_sheet_role(selected.title), f"sheet_name={selected.title}")]
        return unlabelled, fields, warnings, evidence
    finally:
        workbook.close()


def _sheet_role(name: str) -> str | None:
    normalized = re.sub(r"[^A-Z]", "", name.upper())
    if normalized in {"SI", "SHIPPINGINSTRUCTION"}:
        return "SI"
    if normalized in {"BL", "BILLOFLADING"}:
        return "BL"
    return None


def _iter_docx_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _docx_records(source: Path) -> tuple[list[str | FieldLine], list[str], list[tuple[str | None, str | None]]]:
    file = Document(source)
    items: list[str | FieldLine] = []
    warnings: list[str] = []
    table_count = 0
    content: list[str] = []
    for block in _iter_docx_blocks(file):
        if isinstance(block, Paragraph):
            value = _plain(block.text).strip()
            if not value:
                continue
            content.append(value)
            label, separator, remainder = value.partition(":")
            if separator and label.strip():
                items.append(FieldLine(label, remainder))
            else:
                items.append(value)
            continue
        table_count += 1
        for row in block.rows:
            cells = row.cells
            if len(cells) != 2:
                warnings.append(f"non_two_cell_table_row:{table_count}")
                items.append(" | ".join(_plain(cell.text).strip() for cell in cells if cell.text.strip()))
                continue
            label = _plain(cells[0].text).strip()
            value = "\n".join(_plain(paragraph.text).strip() for paragraph in cells[1].paragraphs if paragraph.text.strip())
            content.extend([label, value])
            if label:
                items.append(FieldLine(label, value))
            elif value:
                warnings.append(f"blank_label_table_row:{table_count}")
                items.append(value)
    if table_count == 0:
        warnings.append("no_table")
    elif table_count > 1:
        warnings.append(f"multiple_tables:{table_count}")
    role, detail = _role_from_text("\n".join(content))
    return items, warnings, [(role, detail)]


def _read_pdf_layout(source: Path) -> tuple[str, int, bool, list[str]]:
    errors: list[str] = []
    try:
        pdf = pymupdf.open(source)
        try:
            pages = len(pdf)
            if pages == 0:
                raise ConversionError("zero_page_pdf")
            font_present = any(page.get_fonts() for page in pdf)
            fallback = "\n".join(_pymupdf_layout(page) for page in pdf)
        finally:
            pdf.close()
    except Exception as error:
        errors.append(type(error).__name__)
        pages, font_present, fallback = 0, False, ""
    command = shutil.which("pdftotext")
    if command:
        completed = subprocess.run(
            [command, "-layout", str(source), "-"], capture_output=True, text=True, encoding="utf-8", errors="replace", check=False
        )
        if completed.returncode == 0:
            return completed.stdout.replace("\r\n", "\n"), pages, font_present, errors
        errors.append("pdftotext_failed")
    else:
        errors.append("pdftotext_unavailable")
    if pages == 0:
        try:
            logging.getLogger("pypdf").setLevel(logging.ERROR)
            reader = PdfReader(source)
            pages = len(reader.pages)
            fallback = "\n".join(page.extract_text() or "" for page in reader.pages)
            font_present = bool(fallback.strip())
            errors.append("pymupdf_failed_pypdf_recovered")
        except Exception:
            raise ConversionError("corrupt_pdf") from None
    return fallback.replace("\r\n", "\n"), pages, font_present, errors


def _pymupdf_layout(page: pymupdf.Page) -> str:
    """Recover the label/value visual columns when Poppler is not installed."""
    lines: list[str] = []
    for block in page.get_text("blocks", sort=True):
        x0 = block[0]
        values = [line.strip() for line in block[4].splitlines() if line.strip()]
        if not values:
            continue
        if len(values) == 2 and ":" not in values[0]:
            lines.append(f"{values[0]}  {values[1]}")
        elif x0 > 100 and lines:
            lines.extend(f"  {value}" for value in values)
        else:
            lines.extend(values)
    return "\n".join(lines)


def _pdf_layout_records(layout: str) -> list[str | FieldLine]:
    items: list[str | FieldLine] = []
    for raw in layout.splitlines():
        if not raw.strip():
            continue
        if raw[:1].isspace() and items and isinstance(items[-1], FieldLine):
            previous = items[-1]
            items[-1] = FieldLine(previous.label, previous.value + "\n" + raw.strip())
            continue
        match = LAYOUT_FIELD.match(raw.rstrip())
        upper = raw.upper()
        multiple_labels = sum(marker in upper for marker in ("B/L NUMBER", "BOOKING NO", "HS CODE", "FREIGHT", "OC NO")) >= 2
        if match and not multiple_labels:
            items.append(FieldLine(match.group("label"), match.group("value")))
        else:
            items.append(raw.rstrip())
    return items


def _txt_records(raw: str, source_heading: str) -> list[str | FieldLine]:
    """Keep every line while replacing only the canonical document heading when needed."""
    lines = raw.splitlines()
    body = lines[3:] if len(lines) >= 3 and lines[1] == SEPARATOR else lines
    items: list[str | FieldLine] = [] if source_heading in CANONICAL_HEADINGS.values() else [source_heading]
    for line in body:
        if not line.strip():
            continue
        if line.startswith("  ") and items and isinstance(items[-1], FieldLine):
            previous = items[-1]
            items[-1] = FieldLine(previous.label, previous.value + "\n" + line.strip())
            continue
        label, separator, value = line.partition(":")
        if separator and label.strip():
            items.append(FieldLine(label, value))
        else:
            items.append(line)
    return items


def _render_scan_pages(source: Path, document: dict[str, Any], derived_dir: Path) -> list[Path]:
    page_dir = derived_dir / "pages" / document["doc_id"]
    page_dir.mkdir(parents=True, exist_ok=True)
    pdf = pymupdf.open(source)
    try:
        outputs: list[Path] = []
        for number, page in enumerate(pdf, start=1):
            output = page_dir / f"ocr-page-{number}.png"
            if not output.is_file():
                page.get_pixmap(matrix=pymupdf.Matrix(200 / 72, 200 / 72), alpha=False).save(output)
            outputs.append(output)
        return outputs
    finally:
        pdf.close()


def _tesseract_read(page_paths: list[Path]) -> tuple[str | None, str | None]:
    command = shutil.which("tesseract")
    if not command:
        return None, "tesseract_unavailable"
    lines: list[str] = []
    try:
        for page in page_paths:
            completed = subprocess.run([command, str(page), "stdout", "--psm", "6"], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=45, check=False)
            if completed.returncode != 0:
                return None, "tesseract_failed"
            lines.extend(line for line in completed.stdout.splitlines() if line.strip())
    except (OSError, subprocess.TimeoutExpired):
        return None, "tesseract_failed"
    return "\n".join(lines), None


def _gemini_read(page_paths: list[Path], derived_dir: Path) -> tuple[dict[str, Any] | None, str | None]:
    """Optional second reader. It is inert unless the user explicitly enables it with a key."""
    if os.getenv("GEMINI_ENABLED", "false").lower() != "true":
        return None, "gemini_disabled"
    key = os.getenv("GEMINI_API_KEY")
    if not key:
        return None, "gemini_key_missing"
    model = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite")
    digest = hashlib.sha256((model + GEMINI_PROMPT_VERSION).encode("utf-8") + b"".join(page.read_bytes() for page in page_paths)).hexdigest()
    cache_path = derived_dir / "cache" / "gemini" / "ocr" / f"{digest}.json"
    if cache_path.is_file():
        return json.loads(cache_path.read_text(encoding="utf-8")), None
    global _gemini_last_request
    rpm = max(1, int(os.getenv("GEMINI_MAX_RPM", "12")))
    pause = (60 / rpm) - (clock.monotonic() - _gemini_last_request)
    if pause > 0:
        clock.sleep(pause)
    parts: list[dict[str, Any]] = [{"text": GEMINI_SCAN_PROMPT}]
    parts.extend({"inline_data": {"mime_type": "image/png", "data": base64.b64encode(page.read_bytes()).decode("ascii")}} for page in page_paths)
    request = urllib.request.Request(
        f"https://generativelanguage.googleapis.com/v1beta/models/{urllib.parse.quote(model, safe='')}:generateContent",
        data=json.dumps({"contents": [{"parts": parts}], "generationConfig": {"temperature": 0, "responseMimeType": "application/json"}}).encode("utf-8"),
        headers={"Content-Type": "application/json", "x-goog-api-key": key},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            payload = json.loads(response.read().decode("utf-8"))
        text = payload["candidates"][0]["content"]["parts"][0]["text"]
        reading = json.loads(text)
        if not isinstance(reading.get("raw_lines"), list) or not isinstance(reading.get("fields"), dict):
            return None, "gemini_invalid_schema"
    except (KeyError, json.JSONDecodeError, urllib.error.URLError, TimeoutError):
        return None, "gemini_failed"
    _gemini_last_request = clock.monotonic()
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps(reading, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return reading, None


def _scan_conversion(settings: Settings, source: Path, role: str, document: dict[str, Any], pages: int, warnings: list[str], evidence: list[str]) -> Conversion:
    page_paths = _render_scan_pages(source, document, settings.derived_dir)
    tesseract_text, tesseract_error = _tesseract_read(page_paths)
    gemini_reading, gemini_error = _gemini_read(page_paths, settings.derived_dir)
    reader_warnings = [*warnings, *(error for error in (tesseract_error, gemini_error) if error)]
    if not tesseract_text and not gemini_reading:
        meta = _metadata(document, role, evidence, "failed", "pdf_ocr_tesseract", [*reader_warnings, "scan_readers_unavailable"], pages, 0.0)
        meta["uncertain_fields"] = sorted(SEVEN_FIELDS)
        return Conversion(document["doc_id"], "", meta)
    items = _pdf_layout_records(tesseract_text or "\n".join(str(line) for line in gemini_reading["raw_lines"]))
    text = canonical_text_from_items(role, items)
    local_fields = canonical_fields(text)
    gemini_fields = gemini_reading.get("fields", {}) if gemini_reading else {}
    uncertain = sorted(field for field in SEVEN_FIELDS if not gemini_reading or local_fields.get(field, "").strip() != str(gemini_fields.get(field, "")).strip())
    errors = validate_canonical(text)
    status = "ok" if gemini_reading and not uncertain and not errors else "degraded"
    method = "pdf_ocr_dual" if gemini_reading else "pdf_ocr_tesseract"
    meta = _metadata(document, role, evidence, status, method, [*reader_warnings, *errors], pages, 1.0 if status == "ok" else 0.5)
    meta["uncertain_fields"] = uncertain
    return Conversion(document["doc_id"], text, meta)


def _metadata(document: dict[str, Any], role: str, evidence: list[str], status: str, method: str, warnings: list[str], pages: int | None, confidence: float) -> dict[str, Any]:
    source = Path(document["path"])
    return {
        "doc_id": document["doc_id"],
        "source": source.as_posix(),
        "source_sha256": document["sha256"],
        "role": role,
        "role_evidence": evidence,
        "method": method,
        "status": status,
        "warnings": warnings,
        "pages": pages,
        "confidence": confidence,
        "converted_at": datetime.fromtimestamp(document["source_mtime"], timezone.utc).isoformat(timespec="seconds"),
    }


def convert_document(settings: Settings, document: dict[str, Any]) -> Conversion:
    """Convert one input file. Expected input failures always produce failed metadata."""
    try:
        source = attachment_path(settings.data_dir, document["path"])
    except FileNotFoundError:
        role = document.get("role_hint") or "SI"
        missing = {**document, "source_mtime": 0}
        return Conversion(document["doc_id"], "", _metadata(missing, role, ["attachment_missing"], "failed", "unknown", ["missing_attachment"], None, 0.0))
    document = {**document, "source_mtime": source.stat().st_mtime}
    extension = document["ext"].lower()
    try:
        if extension == ".txt":
            raw = source.read_text(encoding="utf-8").replace("\r\n", "\n").replace("\r", "\n")
            role, evidence, warnings = _resolve_role(document, [_role_from_text(raw)])
            source_heading = raw.splitlines()[0].strip() if raw.splitlines() else ""
            items = _txt_records(raw, source_heading)
            text = canonical_text_from_items(role, items)
            errors = validate_canonical(text, require_shipping_fields=not any(item.startswith("wrong_doc_type_candidate") for item in warnings))
            status = "ok" if not errors else "degraded"
            return Conversion(document["doc_id"], text, _metadata(document, role, evidence, status, "txt_copy", [*warnings, *errors], None, 1.0 if status == "ok" else 0.7))
        if extension == ".xlsx":
            unlabelled, fields, warnings, role_evidence = _xlsx_records(source, document)
            role, evidence, role_warnings = _resolve_role(document, role_evidence)
            text = canonical_text(role, unlabelled, fields)
            errors = validate_canonical(text)
            status = "ok" if not errors else "degraded"
            return Conversion(document["doc_id"], text, _metadata(document, role, evidence, status, "xlsx_openpyxl", [*warnings, *role_warnings, *errors], None, 1.0 if status == "ok" else 0.7))
        if extension == ".docx":
            items, warnings, role_evidence = _docx_records(source)
            role, evidence, role_warnings = _resolve_role(document, role_evidence)
            text = canonical_text_from_items(role, _without_redundant_role_title(items, role))
            errors = validate_canonical(text)
            status = "ok" if not errors else "degraded"
            return Conversion(document["doc_id"], text, _metadata(document, role, evidence, status, "docx_python_docx", [*warnings, *role_warnings, *errors], None, 1.0 if status == "ok" else 0.7))
        if extension == ".pdf":
            layout, pages, font_present, warnings = _read_pdf_layout(source)
            role, evidence, role_warnings = _resolve_role(document, [_role_from_text(layout)])
            if not (font_present and layout.strip()):
                return _scan_conversion(settings, source, role, document, pages, [*warnings, *role_warnings], evidence)
            items = _pdf_layout_records(layout)
            text = canonical_text_from_items(role, items)
            errors = validate_canonical(text)
            status = "ok" if not errors else "degraded"
            return Conversion(document["doc_id"], text, _metadata(document, role, evidence, status, "pdf_text_layout", [*warnings, *role_warnings, *errors], pages, 1.0 if status == "ok" else 0.7))
        raise ConversionError(f"unsupported_format:{extension}")
    except Exception as error:
        fallback_role, fallback_evidence = _role_from_name(document)
        role = fallback_role or "SI"
        method = "pdf_text_layout" if extension == ".pdf" else "unknown"
        return Conversion(document["doc_id"], "", _metadata(document, role, fallback_evidence, "failed", method, [str(error) if isinstance(error, ConversionError) else type(error).__name__], None, 0.0))


def _write_conversion(settings: Settings, conversion: Conversion) -> tuple[str, str]:
    text_dir = settings.derived_dir / "text"
    meta_dir = settings.derived_dir / "meta"
    text_dir.mkdir(parents=True, exist_ok=True)
    meta_dir.mkdir(parents=True, exist_ok=True)
    text_path = text_dir / f"{conversion.doc_id}.txt"
    meta_path = meta_dir / f"{conversion.doc_id}.json"
    text_path.write_text(conversion.text, encoding="utf-8", newline="\n")
    meta_path.write_text(json.dumps(conversion.metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    return f"text/{text_path.name}", f"meta/{meta_path.name}"


def convert_all(settings: Settings, document_ids: set[str] | None = None) -> list[Conversion]:
    """Convert changed attachments only; optional worker processes speed a full run."""
    initialize(settings.database_path)
    with database(settings.database_path) as connection:
        rows = [dict(row) for row in connection.execute("SELECT * FROM documents ORDER BY doc_id")]
    if document_ids is not None:
        rows = [row for row in rows if row["doc_id"] in document_ids]
    pending: list[dict] = []
    cached: list[Conversion] = []
    for row in rows:
        meta_path = settings.derived_dir / (row.get("meta_path") or "")
        text_path = settings.derived_dir / (row.get("text_path") or "")
        try:
            metadata = json.loads(meta_path.read_text(encoding="utf-8"))
            if metadata.get("source_sha256") == row["sha256"] and text_path.is_file():
                cached.append(Conversion(row["doc_id"], text_path.read_text(encoding="utf-8"), metadata))
                continue
        except (OSError, json.JSONDecodeError):
            pass
        pending.append(row)
    workers = max(1, int(os.getenv("CONVERT_WORKERS", "1")))
    if workers > 1 and len(pending) > 1:
        with ProcessPoolExecutor(max_workers=min(workers, len(pending))) as pool:
            fresh = list(pool.map(_convert_worker, [(settings, row) for row in pending]))
    else:
        fresh = [convert_document(settings, row) for row in pending]
    results = [*cached, *fresh]
    with database(settings.database_path) as connection:
        for result in results:
            text_path, meta_path = _write_conversion(settings, result)
            connection.execute(
                """UPDATE documents SET role_detected = ?, convert_status = ?, convert_method = ?, text_path = ?, meta_path = ? WHERE doc_id = ?""",
                (result.metadata["role"], result.metadata["status"], result.metadata["method"], text_path, meta_path, result.doc_id),
            )
    return results


def _convert_worker(item: tuple[Settings, dict[str, Any]]) -> Conversion:
    """Pickle-friendly top-level worker for Windows process pools."""
    settings, document = item
    return convert_document(settings, document)


def summary(results: Iterable[Conversion]) -> dict[str, dict[str, int]]:
    report: dict[str, dict[str, int]] = {}
    for result in results:
        extension = Path(result.metadata["source"]).suffix.lower() or "unknown"
        bucket = report.setdefault(extension, {"ok": 0, "degraded": 0, "failed": 0})
        bucket[result.metadata["status"]] += 1
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--doc-id", action="append", help="convert only one document ID (repeatable)")
    args = parser.parse_args()
    settings = settings_from_env()
    ingest(settings)
    results = convert_all(settings, set(args.doc_id) if args.doc_id else None)
    print(json.dumps(summary(results), indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
