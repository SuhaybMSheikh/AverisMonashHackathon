"""Safe, read-only previews for the participant attachment formats."""
from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from openpyxl import load_workbook


MAX_PREVIEW_ROWS = 1_000
ORIGINAL_GRID_ROWS = 15
ORIGINAL_GRID_COLUMNS = 2
PDF_SCALE = 150 / 72


def _text(value: Any) -> str:
    """Return display text without introducing executable markup."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).replace("\x00", "")


def _label_rows(text: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for line in text.splitlines():
        label, separator, value = line.partition(":")
        if separator and label.strip() and len(label.strip()) <= 160:
            rows.append({"label": label.strip(), "value": value.strip()})
    return rows


def _base(document: dict[str, Any]) -> dict[str, Any]:
    return {
        "doc_id": document["doc_id"],
        "filename": Path(document["path"]).name,
        "format": document["ext"].removeprefix("."),
        "size": document["size"],
        "role_hint": document["role_hint"],
    }


def _unreadable(document: dict[str, Any], error: Exception) -> dict[str, Any]:
    format_name = document["ext"].removeprefix(".").upper() or "document"
    detail = f"The {format_name} reader could not open this file ({type(error).__name__})."
    return {
        **_base(document),
        "error": "unreadable",
        "detail": detail or "The file could not be read.",
        "metadata": {"classification": "corrupt"},
    }


def _xlsx_preview(source: Path, document: dict[str, Any]) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        sheet_name = sheet.title
        sheet_count = len(workbook.sheetnames)
        rows = [
            [_text(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True, max_row=MAX_PREVIEW_ROWS, max_col=2)
        ]
    finally:
        workbook.close()

    nonempty_rows = [row for row in rows if any(value.strip() for value in row)]
    fields = [
        {"label": row[0].strip(), "value": row[1].strip()}
        for row in nonempty_rows
        if row[0].strip()
    ]
    return {
        **_base(document),
        "fields": fields,
        "metadata": {"sheet_name": sheet_name, "sheet_count": sheet_count},
        "original": {
            "kind": "grid",
            "rows": nonempty_rows[:ORIGINAL_GRID_ROWS],
            "columns": ["Column A", "Column B"],
        },
    }


def _docx_preview(source: Path, document: dict[str, Any]) -> dict[str, Any]:
    file = Document(source)
    paragraphs = [_text(paragraph.text).strip() for paragraph in file.paragraphs]
    paragraphs = [paragraph for paragraph in paragraphs if paragraph]
    tables = [
        [[_text(cell.text).strip() for cell in row.cells] for row in table.rows]
        for table in file.tables
    ]
    fields = _label_rows("\n".join(paragraphs))
    for table in tables:
        for row in table:
            if len(row) >= 2 and row[0]:
                fields.append({"label": row[0], "value": row[1]})
    return {
        **_base(document),
        "fields": fields,
        "metadata": {"table_count": len(tables), "paragraph_count": len(paragraphs)},
        "original": {"kind": "document", "paragraphs": paragraphs, "tables": tables},
    }


def _render_pdf_pages(source: Path, document: dict[str, Any], derived_dir: Path) -> tuple[int, list[str], str]:
    pdf = pymupdf.open(source)
    try:
        text = "\n".join(page.get_text("text") for page in pdf)
        page_dir = derived_dir / "pages" / document["doc_id"]
        page_dir.mkdir(parents=True, exist_ok=True)
        page_urls: list[str] = []
        for index, page in enumerate(pdf, start=1):
            output = page_dir / f"page-{index}.png"
            if not output.is_file():
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(PDF_SCALE, PDF_SCALE), alpha=False)
                pixmap.save(output)
            page_urls.append(f"/api/documents/{document['doc_id']}/pages/{index}.png")
        return len(pdf), page_urls, text
    finally:
        pdf.close()


def _pdf_preview(source: Path, document: dict[str, Any], derived_dir: Path) -> dict[str, Any]:
    page_count, page_urls, text = _render_pdf_pages(source, document, derived_dir)
    text_layer = bool(text.strip())
    return {
        **_base(document),
        "fields": _label_rows(text) if text_layer else [],
        "metadata": {
            "page_count": page_count,
            "classification": "text_layer" if text_layer else "scan",
        },
        "original": {"kind": "page_images", "pages": page_urls},
    }


def preview_document(source: Path, document: dict[str, Any], derived_dir: Path, canonical: str | None = None) -> dict[str, Any]:
    """Build a JSON-safe preview; malformed source data is an expected outcome."""
    try:
        extension = document["ext"].lower()
        if extension == ".txt":
            content = source.read_text(encoding="utf-8")
            preview = {
                **_base(document),
                "fields": _label_rows(content),
                "metadata": {"line_count": len(content.splitlines())},
                "original": {"kind": "text", "text": content},
            }
        elif extension == ".xlsx":
            preview = _xlsx_preview(source, document)
        elif extension == ".docx":
            preview = _docx_preview(source, document)
        elif extension == ".pdf":
            preview = _pdf_preview(source, document, derived_dir)
        else:
            raise ValueError(f"Unsupported attachment format: {extension}")
        if canonical is not None:
            preview["fields"] = _label_rows(canonical)
            preview["metadata"] = {**preview.get("metadata", {}), "formatted_source": "canonical_text"}
        return preview
    except Exception as error:  # Files in the inbox are untrusted input.
        return _unreadable(document, error)


def rendered_pdf_page(source: Path, document: dict[str, Any], derived_dir: Path, page_number: int) -> Path | None:
    """Ensure a PDF page cache exists and return one page without trusting a URL path."""
    preview = _pdf_preview(source, document, derived_dir)
    if preview.get("error"):
        return None
    page_count = preview["metadata"]["page_count"]
    if not 1 <= page_number <= page_count:
        return None
    output = derived_dir / "pages" / document["doc_id"] / f"page-{page_number}.png"
    return output if output.is_file() else None
